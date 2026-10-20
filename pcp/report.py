"""Weekly status report (markdown) and the Excel control workbook."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from . import finance, raid, resources, schedule
from .data import BUDGET, ISSUES, PROJECT, RESOURCE_START, RESOURCES, RISKS, STAKEHOLDERS, TASKS

FILL = {"red": "F4B6B6", "amber": "F7DFA3", "green": "BFE3BF"}


def rag(slip, budget_pct, high_risks, overdue_reviews):
    if slip > 10 or budget_pct > 0.10 or high_risks >= 2:
        return "red"
    if slip > 0 or budget_pct > 0.0 or overdue_reviews:
        return "amber"
    return "green"


def build():
    rows, base_finish = schedule.plan(TASKS, PROJECT["start"])
    fc, slip = schedule.forecast(rows, PROJECT["status_date"], TASKS)
    fin_rows = finance.lines(BUDGET)
    fin = finance.totals(fin_rows)
    rk = raid.risks(RISKS, PROJECT["status_date"])
    iss = raid.issues(ISSUES, PROJECT["status_date"])
    res = resources.plan(RESOURCES, RESOURCE_START)
    high = [r for r in rk if r["rating"] == "high"]
    overdue = [r for r in rk if r["overdue"]]
    status = rag(slip, fin["variance_pct"], len(high), len(overdue))
    return dict(rows=rows, base_finish=base_finish, forecast=fc, slip=slip, fin_rows=fin_rows, fin=fin, risks=rk,
                issues=iss, res=res, status=status, high=high, overdue=overdue,
                late=schedule.late_tasks(rows, PROJECT["status_date"]))


def money(v):
    return f"{v:,.0f}".replace(",", ".")


def markdown(b):
    p, f = PROJECT, b["fin"]
    open_issues = [i for i in b["issues"] if i["status"] == "open"]
    over_alloc = [r for r in b["res"] if r["over"]]
    L = [f"# Status report: {p['name']}", "", f"Status date {p['status_date']}. Overall status: **{b['status'].upper()}**.", "",
         "## Summary",
         f"* Schedule: baseline finish {b['base_finish']}, forecast finish {max(b['forecast'].values())}, a slip of {b['slip']} working days.",
         f"* Budget: forecast at completion {money(f['eac'])} {p['currency']} against {money(f['planned'])}, {f['variance_pct']*100:.1f} percent over. {f['spent_pct']*100:.0f} percent of the budget is spent.",
         f"* Risks: {len(b['high'])} high, {len(b['overdue'])} with an overdue review. Open issues: {len(open_issues)}.", "",
         "## Late tasks (should be finished, are not)"]
    L += [f"* {r['id']} {r['name']}: due {r['finish']}, {r['pct']} percent done, owner {r['owner']}" for r in b["late"]] or ["* none"]
    L += ["", "## Decisions needed this week"]
    L += [f"* {i['id']} {i['text']} (open {i['age_days']} days, owner {i['owner']})" for i in open_issues if i["stale"]] or ["* none"]
    L += ["", "## Top risks"]
    L += [f"* {r['id']} {r['text']}: score {r['score']} ({r['rating']}), owner {r['owner']}, mitigation: {r['mitigation']}" + (" (review overdue)" if r["overdue"] else "") for r in b["risks"][:3]]
    L += ["", "## Resources"]
    L += [f"* {r['person']} is over 100 percent in the week of {w} ({pct} percent)" for r in over_alloc for w, pct in r["over"]] or ["* no over-allocation"]
    return "\n".join(L) + "\n"


def workbook(b, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan"
    ws.append(["Id", "Task", "Owner", "Days", "Depends on", "Start", "Finish", "Slack", "Critical", "Percent", "Forecast finish"])
    for r in b["rows"].values():
        ws.append([r["id"], r["name"], r["owner"], r["duration"], ", ".join(r["deps"]), r["start"], r["finish"], r["slack"],
                   "yes" if r["critical"] else "", r["pct"], b["forecast"][r["id"]]])
    ws = wb.create_sheet("Budget")
    ws.append(["Line", "Planned", "Actual", "Estimate to complete", "Forecast", "Variance", "Variance %", "Flag"])
    for r in b["fin_rows"]:
        ws.append([r["line"], r["planned"], r["actual"], r["etc"], r["eac"], r["variance"], round(r["variance_pct"], 3), r["flag"]])
    t = b["fin"]
    ws.append(["Total", t["planned"], t["actual"], "", t["eac"], t["variance"], round(t["variance_pct"], 3), ""])
    ws = wb.create_sheet("Risks")
    ws.append(["Id", "Risk", "P", "I", "Score", "Rating", "Owner", "Mitigation", "Last review", "Review due", "Overdue"])
    for r in b["risks"]:
        ws.append([r["id"], r["text"], r["probability"], r["impact"], r["score"], r["rating"], r["owner"], r["mitigation"], r["last_review"], r["review_due"], "yes" if r["overdue"] else ""])
    ws = wb.create_sheet("Issues")
    ws.append(["Id", "Issue", "Raised", "Owner", "Status", "Age (days)", "Stale"])
    for i in b["issues"]:
        ws.append([i["id"], i["text"], i["raised"], i["owner"], i["status"], i["age_days"], "yes" if i["stale"] else ""])
    ws = wb.create_sheet("Resources")
    ws.append(["Person", "Role"] + [str(RESOURCE_START.fromordinal(RESOURCE_START.toordinal() + 7 * i)) for i in range(6)])
    for r in b["res"]:
        ws.append([r["person"], r["role"]] + r["weeks"])
    ws = wb.create_sheet("Stakeholders")
    ws.append(["Stakeholder", "Interest", "Influence", "How often", "Channel"])
    for s in STAKEHOLDERS:
        ws.append(list(s))
    for sheet in wb.worksheets:
        for c in sheet[1]:
            c.font = Font(bold=True)
    wb["Plan"]["A1"].fill = PatternFill("solid", fgColor=FILL[b["status"]])
    wb.save(path)


def write(outdir):
    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    b = build()
    (out / "status_report.md").write_text(markdown(b))
    workbook(b, out / "control_pack.xlsx")
    return b
