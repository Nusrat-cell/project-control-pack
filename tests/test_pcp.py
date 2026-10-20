from datetime import date

import pytest

from pcp import finance, raid, report, resources, schedule
from pcp.data import BUDGET, PROJECT, RESOURCES, RESOURCE_START, RISKS, TASKS


def test_workdays_skip_weekends():
    assert schedule.add_workdays(date(2026, 10, 9), 2) == date(2026, 10, 12)
    assert schedule.add_workdays(date(2026, 10, 10), 1) == date(2026, 10, 12)


def test_plan_dates_follow_dependencies():
    rows, finish = schedule.plan(TASKS, PROJECT["start"])
    assert rows["T5"]["start"] > rows["T3"]["finish"] and rows["T5"]["start"] > rows["T4"]["finish"]
    assert finish == date(2026, 11, 5)


def test_critical_path_and_slack():
    rows, _ = schedule.plan(TASKS, PROJECT["start"])
    assert [t for t in rows if rows[t]["critical"]] == ["T1", "T2", "T3", "T5", "T6", "T8", "T9", "T10"]
    assert rows["T4"]["slack"] == 3 and rows["T7"]["slack"] == 2


def test_cycle_and_unknown_dependency_are_rejected():
    cyc = [("A", "a", "x", 1, ["B"], 0), ("B", "b", "x", 1, ["A"], 0)]
    with pytest.raises(schedule.PlanError):
        schedule.plan(cyc, date(2026, 1, 5))
    with pytest.raises(schedule.PlanError):
        schedule.plan([("A", "a", "x", 1, ["Z"], 0)], date(2026, 1, 5))


def test_forecast_slip_is_measured_in_working_days():
    rows, _ = schedule.plan(TASKS, PROJECT["start"])
    fc, slip = schedule.forecast(rows, PROJECT["status_date"], TASKS)
    assert fc["T3"] == date(2026, 10, 20) and slip == 12


def test_late_tasks_found():
    rows, _ = schedule.plan(TASKS, PROJECT["start"])
    assert [r["id"] for r in schedule.late_tasks(rows, PROJECT["status_date"])] == ["T3", "T4", "T5", "T7"]


def test_budget_variance_and_flags():
    rows = {r["line"]: r for r in finance.lines(BUDGET)}
    assert rows["Engineering"]["eac"] == 217000.0 and rows["Engineering"]["flag"] == "over"
    assert rows["Discovery and mapping"]["flag"] == "watch"
    t = finance.totals(list(rows.values()))
    assert t["planned"] == 480000.0 and t["eac"] == 536500.0


def test_risk_score_rating_and_overdue_review():
    rk = {r["id"]: r for r in raid.risks(RISKS, PROJECT["status_date"])}
    assert rk["R1"]["score"] == 16 and rk["R1"]["rating"] == "high"
    assert rk["R2"]["overdue"] and rk["R4"]["overdue"] and not rk["R1"]["overdue"]


def test_overallocation_found():
    res = {r["person"]: r for r in resources.plan(RESOURCES, RESOURCE_START)}
    assert res["Engineer"]["over"] == [(date(2026, 11, 2), 120)] and not res["Analyst"]["over"]


def test_rag_rules():
    assert report.rag(0, 0.0, 0, 0) == "green"
    assert report.rag(2, 0.0, 0, 0) == "amber"
    assert report.rag(12, 0.0, 0, 0) == "red"


def test_report_and_workbook_are_written(tmp_path):
    b = report.write(tmp_path)
    text = (tmp_path / "status_report.md").read_text()
    assert b["status"] == "red" and "slip of 12 working days" in text
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "control_pack.xlsx")
    assert wb.sheetnames == ["Plan", "Budget", "Risks", "Issues", "Resources", "Stakeholders"]
